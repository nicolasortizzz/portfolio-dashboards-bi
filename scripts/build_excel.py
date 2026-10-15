#!/usr/bin/env python3
"""Construye un libro Excel por dashboard con tablas formateadas y gráficos nativos."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from fx import excel_ars_format, excel_usd_format, format_kpi_ars_usd, load_exchange_config

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
EXCEL = ROOT / "excel"

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(bold=True, size=14, color="1F4E79")
KPI_FONT = Font(bold=True, size=12)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_table(ws, start_row: int, start_col: int, n_rows: int, n_cols: int) -> None:
    for r in range(start_row, start_row + n_rows):
        for c in range(start_col, start_col + n_cols):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            if r == start_row:
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c > start_col:
                cell.alignment = Alignment(horizontal="right")


def write_df(ws, df: pd.DataFrame, start_row: int = 1, start_col: int = 1) -> tuple[int, int]:
    cols = list(df.columns)
    for j, col in enumerate(cols):
        ws.cell(row=start_row, column=start_col + j, value=col)
    for i, row in enumerate(df.itertuples(index=False), start=1):
        for j, val in enumerate(row):
            ws.cell(row=start_row + i, column=start_col + j, value=val)
    style_table(ws, start_row, start_col, len(df) + 1, len(cols))
    return start_row + len(df), len(cols)


def apply_number_formats(ws, col_names: list[str], formats: dict[str, str], header_row: int = 1) -> None:
    name_to_idx = {name: i + 1 for i, name in enumerate(col_names)}
    for row in ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row):
        for col_name, fmt in formats.items():
            idx = name_to_idx.get(col_name)
            if idx and idx <= len(row):
                cell = row[idx - 1]
                if isinstance(cell.value, (int, float)):
                    cell.number_format = fmt


def autosize(ws, max_col: int) -> None:
    for c in range(1, max_col + 1):
        letter = get_column_letter(c)
        ws.column_dimensions[letter].width = 16


def kpi_block(ws, kpis: list[tuple[str, str]], start_row: int = 3) -> None:
    ws["A1"] = "Resumen de KPIs"
    ws["A1"].font = TITLE_FONT
    row = start_row
    for label, value in kpis:
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row, column=2, value=value).font = KPI_FONT
        row += 1


def build_financiero() -> Path:
    rate = load_exchange_config()["ars_per_usd"]
    pl = pd.read_csv(DATA / "financiero_estado_resultados.csv")
    cf = pd.read_csv(DATA / "financiero_flujo_caja.csv")
    bvr = pd.read_csv(DATA / "financiero_presupuesto_vs_real.csv")

    wb = Workbook()
    ws_kpi = wb.active
    ws_kpi.title = "KPIs"
    last = pl.iloc[-1]
    last_cf = cf.iloc[-1]
    kpi_block(
        ws_kpi,
        [
            ("Ingresos último mes", format_kpi_ars_usd(last["ingresos_ars"], rate)),
            ("EBITDA último mes", format_kpi_ars_usd(last["ebitda_ars"], rate)),
            ("Resultado neto último mes", format_kpi_ars_usd(last["resultado_neto_ars"], rate)),
            ("Margen EBITDA (último mes)", f"{last['ebitda_ars'] / last['ingresos_ars'] * 100:.1f}%"),
            ("Flujo neto último mes", format_kpi_ars_usd(last_cf["flujo_neto_ars"], rate)),
            ("Saldo de caja", format_kpi_ars_usd(last_cf["saldo_caja_ars"], rate)),
        ],
    )

    ws_pl = wb.create_sheet("Estado_Resultados")
    end_row, ncols = write_df(ws_pl, pl)
    apply_number_formats(
        ws_pl,
        list(pl.columns),
        {c: excel_ars_format() for c in pl.columns if c.endswith("_ars")}
        | {c: excel_usd_format() for c in pl.columns if c.endswith("_usd")},
    )
    ws_pl.freeze_panes = "A2"
    autosize(ws_pl, ncols)

    chart = LineChart()
    chart.title = "Ingresos y EBITDA (ARS)"
    chart.y_axis.title = "ARS"
    meses = Reference(ws_pl, min_col=2, min_row=1, max_row=end_row)
    ing = Reference(ws_pl, min_col=3, min_row=1, max_row=end_row)
    ebitda = Reference(ws_pl, min_col=13, min_row=1, max_row=end_row)
    chart.add_data(ing, titles_from_data=True)
    chart.add_data(ebitda, titles_from_data=True)
    chart.set_categories(meses)
    ws_pl.add_chart(chart, f"E{end_row + 2}")

    ws_cf = wb.create_sheet("Flujo_Caja")
    end_cf, nc_cf = write_df(ws_cf, cf)
    apply_number_formats(
        ws_cf,
        list(cf.columns),
        {c: excel_ars_format() for c in cf.columns if c.endswith("_ars")}
        | {c: excel_usd_format() for c in cf.columns if c.endswith("_usd")},
    )
    ws_cf.freeze_panes = "A2"
    autosize(ws_cf, nc_cf)

    bar = BarChart()
    bar.type = "col"
    bar.title = "Flujo neto mensual (ARS)"
    cats = Reference(ws_cf, min_col=2, min_row=2, max_row=end_cf)
    data = Reference(ws_cf, min_col=list(cf.columns).index("flujo_neto_ars") + 1, min_row=1, max_row=end_cf)
    bar.add_data(data, titles_from_data=True)
    bar.set_categories(cats)
    ws_cf.add_chart(bar, "K2")

    ws_b = wb.create_sheet("Presupuesto_vs_Real")
    end_b, nc_b = write_df(ws_b, bvr)
    apply_number_formats(
        ws_b,
        list(bvr.columns),
        {c: excel_ars_format() for c in bvr.columns if "ars" in c}
        | {c: "0.0%" if c == "variacion_pct" else excel_usd_format() for c in bvr.columns if c.endswith("_usd")},
    )
    if "variacion_pct" in bvr.columns:
        idx = list(bvr.columns).index("variacion_pct") + 1
        for r in range(2, end_b + 1):
            ws_b.cell(row=r, column=idx).number_format = "0.00"
    ws_b.freeze_panes = "A2"
    autosize(ws_b, nc_b)

    out = EXCEL / "dashboard_financiero.xlsx"
    EXCEL.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def build_ventas() -> Path:
    rate = load_exchange_config()["ars_per_usd"]
    mensual = pd.read_csv(DATA / "ventas_mensual.csv")
    canal = pd.read_csv(DATA / "ventas_por_canal.csv")
    region = pd.read_csv(DATA / "ventas_por_region.csv")
    clientes = pd.read_csv(DATA / "ventas_top_clientes.csv")

    wb = Workbook()
    ws_kpi = wb.active
    ws_kpi.title = "KPIs"
    kpi_block(
        ws_kpi,
        [
            ("Ventas último mes", format_kpi_ars_usd(mensual.iloc[-1]["ventas_netas_ars"], rate)),
            ("Ventas acumuladas 24m", format_kpi_ars_usd(mensual["ventas_netas_ars"].sum(), rate)),
            ("Unidades último mes", f"{mensual.iloc[-1]['unidades']:,}".replace(",", ".")),
            ("Top cliente (ficticio)", clientes.iloc[0]["cliente"]),
            ("Ventas top cliente 12m", format_kpi_ars_usd(clientes.iloc[0]["ventas_12m_ars"], rate)),
        ],
    )

    ws_m = wb.create_sheet("Ventas_Mensual")
    end_m, nc_m = write_df(ws_m, mensual)
    apply_number_formats(
        ws_m,
        list(mensual.columns),
        {c: excel_ars_format() for c in mensual.columns if c.endswith("_ars")}
        | {c: excel_usd_format() for c in mensual.columns if c.endswith("_usd")},
    )
    ws_m.freeze_panes = "A2"
    autosize(ws_m, nc_m)

    line = LineChart()
    line.title = "Ventas netas mensuales (ARS)"
    cats = Reference(ws_m, min_col=2, min_row=2, max_row=end_m)
    data = Reference(ws_m, min_col=3, min_row=1, max_row=end_m)
    line.add_data(data, titles_from_data=True)
    line.set_categories(cats)
    ws_m.add_chart(line, "H2")

    last_month = mensual.iloc[-1]["mes"]
    canal_last = canal[canal["mes"] == last_month]
    ws_c = wb.create_sheet("Por_Canal")
    end_c, nc_c = write_df(ws_c, canal_last)
    apply_number_formats(ws_c, list(canal_last.columns), {c: excel_ars_format() for c in canal_last.columns if c.endswith("_ars")})
    ws_c.freeze_panes = "A2"

    pie = PieChart()
    pie.title = f"Mix por canal — {last_month}"
    labels = Reference(ws_c, min_col=4, min_row=2, max_row=end_c)
    data = Reference(ws_c, min_col=5, min_row=2, max_row=end_c)
    pie.add_data(data)
    pie.set_categories(labels)
    ws_c.add_chart(pie, "H2")

    reg_last = region[region["mes"] == last_month]
    ws_r = wb.create_sheet("Por_Region")
    end_r, nc_r = write_df(ws_r, reg_last)
    apply_number_formats(ws_r, list(reg_last.columns), {c: excel_ars_format() for c in reg_last.columns if c.endswith("_ars")})
    ws_r.freeze_panes = "A2"

    bar = BarChart()
    bar.title = f"Ventas por región — {last_month}"
    cats = Reference(ws_r, min_col=4, min_row=2, max_row=end_r)
    data = Reference(ws_r, min_col=5, min_row=1, max_row=end_r)
    bar.add_data(data, titles_from_data=True)
    bar.set_categories(cats)
    ws_r.add_chart(bar, "H2")

    ws_cli = wb.create_sheet("Top_Clientes")
    end_cli, _ = write_df(ws_cli, clientes.head(10))
    apply_number_formats(
        ws_cli,
        list(clientes.columns),
        {c: excel_ars_format() for c in clientes.columns if c.endswith("_ars")}
        | {c: excel_usd_format() for c in clientes.columns if c.endswith("_usd")},
    )
    ws_cli.freeze_panes = "A2"

    bar2 = BarChart()
    bar2.title = "Top 10 clientes (12m ARS)"
    cats = Reference(ws_cli, min_col=3, min_row=2, max_row=min(11, end_cli))
    data = Reference(ws_cli, min_col=4, min_row=1, max_row=min(11, end_cli))
    bar2.add_data(data, titles_from_data=True)
    bar2.set_categories(cats)
    ws_cli.add_chart(bar2, "G2")

    out = EXCEL / "dashboard_ventas.xlsx"
    EXCEL.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def build_costos() -> Path:
    rate = load_exchange_config()["ars_per_usd"]
    fv = pd.read_csv(DATA / "costos_fijo_variable.csv")
    cc = pd.read_csv(DATA / "costos_por_centro.csv")
    unit = pd.read_csv(DATA / "costos_unitarios.csv")
    var = pd.read_csv(DATA / "costos_variaciones.csv")

    wb = Workbook()
    ws_kpi = wb.active
    ws_kpi.title = "KPIs"
    last = fv.iloc[-1]
    kpi_block(
        ws_kpi,
        [
            ("Costo total último mes", format_kpi_ars_usd(last["costo_total_ars"], rate)),
            ("Costos fijos", format_kpi_ars_usd(last["costos_fijos_ars"], rate)),
            ("Costos variables", format_kpi_ars_usd(last["costos_variables_ars"], rate)),
            ("Unidades producidas", f"{last['unidades_producidas']:,}".replace(",", ".")),
            ("Variación costo último mes", f"{var.iloc[-1]['variacion_pct']:.2f}%".replace(".", ",")),
        ],
    )

    ws_fv = wb.create_sheet("Fijo_Variable")
    end_fv, nc_fv = write_df(ws_fv, fv)
    apply_number_formats(
        ws_fv,
        list(fv.columns),
        {c: excel_ars_format() for c in fv.columns if c.endswith("_ars")}
        | {c: excel_usd_format() for c in fv.columns if c.endswith("_usd")},
    )
    ws_fv.freeze_panes = "A2"
    autosize(ws_fv, nc_fv)

    line = LineChart()
    line.title = "Evolución costos fijos vs variables (ARS)"
    cats = Reference(ws_fv, min_col=2, min_row=2, max_row=end_fv)
    fijos = Reference(ws_fv, min_col=3, min_row=1, max_row=end_fv)
    vars_ = Reference(ws_fv, min_col=5, min_row=1, max_row=end_fv)
    line.add_data(fijos, titles_from_data=True)
    line.add_data(vars_, titles_from_data=True)
    line.set_categories(cats)
    ws_fv.add_chart(line, "J2")

    last_month = fv.iloc[-1]["mes"]
    cc_last = cc[cc["mes"] == last_month]
    ws_cc = wb.create_sheet("Por_Centro")
    end_cc, _ = write_df(ws_cc, cc_last)
    apply_number_formats(
        ws_cc,
        list(cc_last.columns),
        {c: excel_ars_format() for c in cc_last.columns if c.endswith("_ars")}
        | {c: excel_usd_format() for c in cc_last.columns if c.endswith("_usd")},
    )
    ws_cc.freeze_panes = "A2"

    pie = PieChart()
    pie.title = f"Costos por centro — {last_month}"
    labels = Reference(ws_cc, min_col=4, min_row=2, max_row=end_cc)
    data = Reference(ws_cc, min_col=5, min_row=2, max_row=end_cc)
    pie.add_data(data)
    pie.set_categories(labels)
    ws_cc.add_chart(pie, "H2")

    unit_last = unit[unit["mes"] == last_month]
    ws_u = wb.create_sheet("Costo_Unitario")
    end_u, _ = write_df(ws_u, unit_last)
    apply_number_formats(
        ws_u,
        list(unit_last.columns),
        {c: excel_ars_format() for c in unit_last.columns if "unitario" in c and c.endswith("_ars")}
        | {c: excel_usd_format() for c in unit_last.columns if "unitario" in c and c.endswith("_usd")},
    )
    ws_u.freeze_panes = "A2"

    bar = BarChart()
    bar.title = f"Costo unitario por producto — {last_month}"
    cats = Reference(ws_u, min_col=4, min_row=2, max_row=end_u)
    data = Reference(ws_u, min_col=5, min_row=1, max_row=end_u)
    bar.add_data(data, titles_from_data=True)
    bar.set_categories(cats)
    ws_u.add_chart(bar, "H2")

    ws_v = wb.create_sheet("Variaciones")
    end_v, _ = write_df(ws_v, var)
    apply_number_formats(
        ws_v,
        list(var.columns),
        {c: excel_ars_format() for c in var.columns if c.endswith("_ars")}
        | {c: excel_usd_format() for c in var.columns if c.endswith("_usd")},
    )
    ws_v.freeze_panes = "A2"

    line2 = LineChart()
    line2.title = "Variación de costos (%)"
    cats = Reference(ws_v, min_col=2, min_row=2, max_row=end_v)
    data = Reference(ws_v, min_col=7, min_row=1, max_row=end_v)
    line2.add_data(data, titles_from_data=True)
    line2.set_categories(cats)
    ws_v.add_chart(line2, "I2")

    out = EXCEL / "dashboard_costos.xlsx"
    EXCEL.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def main() -> None:
    paths = [build_financiero(), build_ventas(), build_costos()]
    for p in paths:
        print(f"Excel generado: {p}")


if __name__ == "__main__":
    main()
